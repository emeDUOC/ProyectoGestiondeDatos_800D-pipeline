"""
Genera el resumen del pipeline en formato Word (.docx) y texto plano (.txt).
Salida: data/reports/pipeline_summary.docx y data/reports/pipeline_summary.txt
"""

from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT_DIR = Path("data/reports")

# Datos del pipeline adaptados al negocio de Taxis de Chicago
STAGES = [
    (
        "Stage 1 — Limpieza (src/clean.py)",
        (
            "Detecta y elimina de forma automática columnas vacías o de relleno que contengan 'zero'. "
            "Traduce y homologa los nombres de las variables al español técnico para cumplir con el esquema "
            "de base de datos (id_taxi, fecha_viaje, duracion_segundos, distancia_millas, total_pago). "
            "Elimina duplicados operativos primarios e imputa inteligentemente valores nulos mediante el cálculo "
            "de la mediana. Genera la columna derivada 'categoria_viaje' segmentando por rangos de millas."
        ),
    ),
    (
        "Stage 2 — Validación semántica (src/validate.py)",
        (
            "Aplica un motor de control de calidad (Data Quality) basado en reglas de negocio duras: "
            "integridad de llaves, duraciones positivas, distancias válidas y un umbral máximo de facturación "
            "de 1000 USD por viaje. Segrega los registros conformes de las anomalías semánticas, "
            "almacenando de forma aislada los rechazos y escribiendo un informe sintético de auditoría "
            "en data/reports/validation_report.txt."
        ),
    ),
    (
        "Stage 3 — Carga a PostgreSQL (src/load.py)",
        (
            "Garantiza el plano DDL inicial ejecutando el script create_table.sql con 5 restricciones avanzadas "
            "(CHECK constraints y UNIQUE keys). Realiza una inserción controlada registro a registro en la base de datos "
            "PostgreSQL, aislando en un archivo secundario cualquier conflicto operativo por duplicidad de registros "
            "para salvaguardar la integridad transaccional."
        ),
    ),
]

FILES_TABLE = [
    ("src/clean.py",           "Stage 1: Extracción, limpieza y feature engineering"),
    ("src/validate.py",        "Stage 2: Validación semántica y control de calidad"),
    ("src/load.py",            "Stage 3: Carga transaccional y control de excepciones"),
    ("main.py",                "Orquestador avanzado por despachador con parámetro --stage"),
    ("sql/create_table.sql",   "DDL de la tabla viajes_taxis con 5 restricciones de negocio"),
    ("docker-compose.yml",     "Infraestructura aislada: PostgreSQL 16 + pgAdmin 4"),
    ("Dockerfile",             "Imagen optimizada del pipeline (Python 3.12-slim + uv)"),
    (".env",                   "Variables de entorno y secretos de conexión local"),
    ("README.md",              "Documentación operativa del sistema DataOps"),
]

COMMANDS = [
    "# Levantar infraestructura y base de datos en segundo plano",
    "docker compose up postgres pgadmin -d",
    "",
    "# Instalar dependencias en el entorno local",
    "pip install pandas psycopg2-binary python-dotenv python-docx sqlalchemy",
    "",
    "# Ejecutar el pipeline orquestado completo",
    "python main.py",
    "",
    "# Ejecutar etapas individuales para demostrar modularidad",
    "python main.py --stage clean",
    "python main.py --stage validate",
    "python main.py --stage load",
    "",
    "# Generar reportes corporativos automatizados",
    "python export_summary.py",
]


# ──────────────────────────────────────────────
# Word (.docx)
# ──────────────────────────────────────────────

def _heading(doc: Document, text: str, level: int) -> None:
    doc.add_heading(text, level=level)


def _add_table(doc: Document, headers: list[str], rows: list[tuple]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        run = hdr[i].paragraphs[0].runs[0]
        run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = val


def build_word(out_path: Path) -> None:
    doc = Document()

    # Title
    title = doc.add_heading("Resumen del Pipeline — Control de Operaciones de Taxis", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Archivos creados
    _heading(doc, "Arquitectura de Componentes Creados", level=1)
    _add_table(doc, ["Archivo / Ruta", "Rol de Ingeniería de Datos"], FILES_TABLE)
    doc.add_paragraph()

    # Etapas
    _heading(doc, "Descripción Funcional de Etapas (ETL)", level=1)
    for stage_title, stage_body in STAGES:
        _heading(doc, stage_title, level=2)
        doc.add_paragraph(stage_body)

    # Comandos de ejecución
    _heading(doc, "Manual de Ejecución y Secuencia Operativa", level=1)
    for line in COMMANDS:
        p = doc.add_paragraph()
        run = p.add_run(line)
        run.font.name = "Courier New"
        run.font.size = Pt(9)
        if line.startswith("#"):
            run.font.color.rgb = RGBColor(0x6A, 0x6A, 0x6A)

    # pgAdmin
    _heading(doc, "Administración de Datos (pgAdmin)", level=1)
    doc.add_paragraph(
        "La interfaz web está disponible de forma nativa en http://localhost:8080 con las credenciales "
        "mapeadas en el archivo de configuración .env. Al dar de alta el servidor en pgAdmin, se debe utilizar "
        "como host la cadena 'postgres', permitiendo la resolución de nombres dentro de la red Docker bridge."
    )

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    print(f"Word guardado exitosamente en: {out_path}")


# ──────────────────────────────────────────────
# Texto plano (.txt)
# ──────────────────────────────────────────────

def build_txt(out_path: Path) -> None:
    sep = "=" * 72
    lines: list[str] = []

    lines += [
        sep,
        "   RESUMEN DEL PIPELINE — CONTROL DE OPERACIONES DE TAXIS (CHICAGO)",
        sep,
        "",
    ]

    # Archivos creados
    lines += ["ARQUITECTURA DE COMPONENTES CREADOS", "-" * 40]
    col_w = max(len(f) for f, _ in FILES_TABLE) + 2
    for fname, rol in FILES_TABLE:
        lines.append(f"  {fname:<{col_w}}{rol}")
    lines.append("")

    # Etapas
    lines += ["DESCRIPCIÓN FUNCIONAL DE ETAPAS (ETL)", "-" * 40]
    for stage_title, stage_body in STAGES:
        lines.append(f"\n  {stage_title}")
        words = stage_body.split()
        current_line = "    "
        for word in words:
            if len(current_line) + len(word) + 1 > 72:
                lines.append(current_line.rstrip())
                current_line = "    " + word + " "
            else:
                current_line += word + " "
        if current_line.strip():
            lines.append(current_line.rstrip())
    lines.append("")

    # Comandos
    lines += ["MANUAL DE EJECUCIÓN Y SECUENCIA OPERATIVA", "-" * 40]
    for cmd in COMMANDS:
        lines.append(f"  {cmd}" if cmd else "")
    lines.append("")

    # pgAdmin
    lines += [
        "ADMINISTRACIÓN DE DATOS (PGADMIN)",
        "-" * 40,
        "  Disponible en el puerto web http://localhost:8080 según credenciales de .env.",
        "  Definir el parámetro host como 'postgres' al enlazar el servidor.",
        "",
        sep,
    ]

    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text("\n".join(lines), encoding="utf-8")
    print(f"Texto plano guardado exitosamente en: {out_path}")


if __name__ == "__main__":
    build_word(OUT_DIR / "pipeline_summary.docx")
    build_txt(OUT_DIR / "pipeline_summary.txt")